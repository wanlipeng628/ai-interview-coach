"""Tests for resume upload endpoint."""

from datetime import datetime
from io import BytesIO

import pytest
from fastapi.testclient import TestClient

from app.api.routes.resume import get_resume_service
from app.application.resume.resume_dto import ResumeProfileResponse
from app.application.resume.resume_service import ResumeService
from app.infrastructure.repositories.resume_repository import ResumeRepository
from app.main import app


class MockResumeRepository(ResumeRepository):
    """In-memory resume repository for testing."""

    def __init__(self) -> None:
        self._profiles: dict[int, ResumeProfileResponse] = {}
        self._next_id = 1

    def get_default(self, user_id: int) -> ResumeProfileResponse | None:
        return self._profiles.get(user_id)

    def upsert_default(
        self,
        user_id: int,
        title: str,
        content: str,
    ) -> ResumeProfileResponse:
        existing = self._profiles.get(user_id)
        now = datetime.now().isoformat()
        summary = content[:160] if content else None
        if existing is None:
            profile = ResumeProfileResponse(
                id=self._next_id,
                title=title,
                content=content,
                summary=summary,
                is_default=True,
                update_time=now,
            )
            self._next_id += 1
        else:
            profile = ResumeProfileResponse(
                id=existing.id,
                title=title,
                content=content,
                summary=summary,
                is_default=True,
                update_time=now,
            )
        self._profiles[user_id] = profile
        return profile


def _get_test_service() -> ResumeService:
    return ResumeService(MockResumeRepository())


@pytest.fixture
def client() -> TestClient:
    """Test client with mocked resume repository."""
    app.dependency_overrides[get_resume_service] = _get_test_service
    yield TestClient(app)
    app.dependency_overrides.pop(get_resume_service, None)


class TestUploadResumeTxt:
    def test_upload_txt_success(self, client: TestClient) -> None:
        content = "张三\n后端工程师\n5年Python开发经验"
        response = client.post(
            "/api/resume/upload",
            files={"file": ("resume.txt", content.encode("utf-8"), "text/plain")},
        )
        assert response.status_code == 201
        body = response.json()
        assert body["title"] == "resume"
        assert "张三" in body["content"]
        assert "后端工程师" in body["content"]
        assert body["is_default"] is True
        assert body["id"] > 0

    def test_upload_md_success(self, client: TestClient) -> None:
        content = "# 李四\n## 技能\n- Java\n- Spring Boot"
        response = client.post(
            "/api/resume/upload",
            files={"file": ("resume.md", content.encode("utf-8"), "text/markdown")},
        )
        assert response.status_code == 201
        body = response.json()
        assert body["title"] == "resume"
        assert "李四" in body["content"]
        assert "Java" in body["content"]


class TestUploadResumePdf:
    def test_upload_blank_pdf_returns_400(self, client: TestClient) -> None:
        from pypdf import PdfWriter

        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        buf = BytesIO()
        writer.write(buf)
        pdf_bytes = buf.getvalue()

        response = client.post(
            "/api/resume/upload",
            files={"file": ("resume.pdf", pdf_bytes, "application/pdf")},
        )
        assert response.status_code == 400


class TestUploadResumeDocx:
    def test_upload_docx_success(self, client: TestClient) -> None:
        from docx import Document

        doc = Document()
        doc.add_paragraph("王五")
        doc.add_paragraph("前端工程师")
        doc.add_paragraph("React, Vue, TypeScript")
        buf = BytesIO()
        doc.save(buf)
        docx_bytes = buf.getvalue()

        response = client.post(
            "/api/resume/upload",
            files={
                "file": (
                    "resume.docx",
                    docx_bytes,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        assert response.status_code == 201
        body = response.json()
        assert body["title"] == "resume"
        assert "王五" in body["content"]
        assert "前端工程师" in body["content"]


class TestUploadUnsupportedType:
    def test_unsupported_file_type_returns_400(self, client: TestClient) -> None:
        response = client.post(
            "/api/resume/upload",
            files={"file": ("resume.rtf", b"some rtf content", "application/rtf")},
        )
        assert response.status_code == 400
        body = response.json()
        assert "Unsupported" in body["detail"] or "unsupported" in body["detail"]

    def test_doc_file_not_docx_returns_400(self, client: TestClient) -> None:
        response = client.post(
            "/api/resume/upload",
            files={"file": ("resume.doc", b"old word format", "application/msword")},
        )
        assert response.status_code == 400

    def test_image_file_returns_400(self, client: TestClient) -> None:
        response = client.post(
            "/api/resume/upload",
            files={"file": ("resume.jpg", b"fake image data", "image/jpeg")},
        )
        assert response.status_code == 400


class TestUploadEdgeCases:
    def test_filename_without_extension(self, client: TestClient) -> None:
        response = client.post(
            "/api/resume/upload",
            files={"file": ("resume", b"some text", "text/plain")},
        )
        assert response.status_code == 400

    def test_empty_file_content_txt(self, client: TestClient) -> None:
        response = client.post(
            "/api/resume/upload",
            files={"file": ("empty.txt", b"", "text/plain")},
        )
        assert response.status_code == 400

    def test_whitespace_only_file(self, client: TestClient) -> None:
        response = client.post(
            "/api/resume/upload",
            files={"file": ("blank.txt", b"   \n\n  ", "text/plain")},
        )
        assert response.status_code == 400

    def test_long_filename_truncated(self, client: TestClient) -> None:
        long_name = "a" * 150 + ".txt"
        content = "测试内容"
        response = client.post(
            "/api/resume/upload",
            files={"file": (long_name, content.encode("utf-8"), "text/plain")},
        )
        assert response.status_code == 201
        body = response.json()
        assert len(body["title"]) <= 128


class TestResponseShape:
    def test_upload_response_matches_resume_profile_response(
        self, client: TestClient
    ) -> None:
        content = "测试简历"
        response = client.post(
            "/api/resume/upload",
            files={"file": ("test.txt", content.encode("utf-8"), "text/plain")},
        )
        assert response.status_code == 201
        body = response.json()
        assert set(body.keys()) >= {
            "id",
            "title",
            "content",
            "summary",
            "is_default",
            "update_time",
        }
        assert isinstance(body["id"], int)
        assert isinstance(body["title"], str)
        assert isinstance(body["content"], str)
        assert isinstance(body["is_default"], bool)
        assert isinstance(body["update_time"], str)
